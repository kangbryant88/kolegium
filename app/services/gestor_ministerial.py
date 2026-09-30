"""
Gestor Ministerial - Fase 3 ("El Músculo").

Recibe una plantilla Excel vacía de la Zona Educativa y una instrucción en
lenguaje natural; el Cerebro (ai_ministerial) decide qué filtrar y qué dato
va en cada columna, y aquí se rellena la plantilla con openpyxl conservando
estilos, bordes, celdas combinadas y logos.

Solo viajan a Gemini los encabezados y la instrucción: ningún dato personal.
"""
import io
import re
import unicodedata
from copy import copy
from datetime import date

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.models import Usuario, Rol
from app.services.ai_ministerial import analizar_formato_ministerio
from app.services.ficha_ministerial import nombre_para_mostrar

FILAS_ENCABEZADO = 10          # el encabezado se busca en las primeras 10 filas
LARGO_MAXIMO_ENCABEZADO = 60   # celdas más largas son títulos o advertencias
# Columna de numeración ("N°", "Nro."): no es un dato del trabajador, se rellena 1, 2, 3...
ENCABEZADOS_NUMERACION = {'n', 'no', 'nro', 'num', 'numero', 'item', 'n o'}
PALABRAS_VACIAS = {'de', 'del', 'la', 'las', 'los', 'el', 'y', 'en', 'para', 'con'}


class ErrorGestor(Exception):
    """Error que se le puede mostrar tal cual al administrador."""


def _normalizar(texto):
    texto = unicodedata.normalize('NFKD', str(texto or '')).encode('ascii', 'ignore').decode()
    return ' '.join(re.sub(r'[^a-z0-9]+', ' ', texto.lower()).split())


# ==========================================
# --- LECTURA DE LA PLANTILLA ---
# ==========================================

def detectar_encabezados(hoja):
    """
    Fila de encabezados = la que tenga más celdas con texto entre las primeras
    FILAS_ENCABEZADO. Devuelve (número de fila, {columna: texto}).
    """
    mejor_fila, mejores = None, {}
    for fila in hoja.iter_rows(min_row=1, max_row=FILAS_ENCABEZADO):
        textos = {}
        for celda in fila:
            if isinstance(celda, MergedCell) or not isinstance(celda.value, str):
                continue
            texto = ' '.join(celda.value.split())
            if texto and len(texto) <= LARGO_MAXIMO_ENCABEZADO:
                textos[celda.column] = texto
        if len(textos) > len(mejores):
            mejor_fila, mejores = fila[0].row, textos
    return mejor_fila, mejores


# ==========================================
# --- FILTRO DE TRABAJADORES ---
# ==========================================

def _coincide_cargo(pedido, usuario):
    """
    'Docentes fuera de aula' coincide con cargo 'Docente Fuera De Aula': cada
    palabra pedida (sin plural) debe ser el comienzo de alguna palabra del
    cargo, del tipo de personal o del área de trabajo.
    """
    palabras_pedidas = [p for p in _normalizar(pedido).split() if p not in PALABRAS_VACIAS]
    if not palabras_pedidas:
        return True
    for texto in (usuario.cargo, usuario.tipo_personal, usuario.area_trabajo):
        palabras = _normalizar(texto).split()
        if palabras and all(any(w.startswith(p[:max(4, len(p) - 2)]) for w in palabras)
                            for p in palabras_pedidas):
            return True
    return False


def filtrar_trabajadores(filtros):
    """Usuarios con rol asignado (nunca los 'En Espera'), según los filtros del Cerebro."""
    consulta = (Usuario.query.join(Rol, Usuario.rol_id == Rol.id)
                .filter(Usuario.rol_id.isnot(None), Rol.nombre != 'En Espera'))
    if filtros.get('solo_activos', True):
        consulta = consulta.filter(Usuario.activo.is_(True))
    usuarios = consulta.all()

    cargo = filtros.get('cargo_requerido')
    if cargo:
        usuarios = [u for u in usuarios if _coincide_cargo(cargo, u)]
    return sorted(usuarios, key=lambda u: ((u.apellidos or u.nombre_completo or '').lower(),
                                           (u.nombres or '').lower()))


# ==========================================
# --- INYECCIÓN ---
# ==========================================

def _valor(usuario, atributo):
    if atributo == 'nombre_completo':
        return nombre_para_mostrar(usuario)
    valor = getattr(usuario, atributo)
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)  # horas: 36 y no 36.0
    return valor if valor != '' else None


def inyectar(hoja, fila_encabezado, columnas, columnas_numeracion, usuarios):
    """
    Escribe un trabajador por fila desde la fila siguiente al encabezado.
    Las filas nuevas copian el estilo (bordes, fuente) de la primera fila de
    datos. Si una celda a escribir ya tiene contenido (p. ej. el pie con las
    firmas) se detiene sin escribir nada, para no dañar la plantilla.
    """
    inicio = fila_encabezado + 1
    todas = {**columnas, **{c: None for c in columnas_numeracion}}

    for i in range(len(usuarios)):
        for col in todas:
            celda = hoja.cell(row=inicio + i, column=col)
            # La numeración puede venir ya escrita en la plantilla: se reescribe igual
            ocupada = celda.value not in (None, '') and col not in columnas_numeracion
            if isinstance(celda, MergedCell) or ocupada:
                raise ErrorGestor(
                    f'La plantilla no tiene espacio para {len(usuarios)} trabajadores: la fila {inicio + i} '
                    f'ya tiene contenido en la columna {celda.column_letter}. Deja al menos '
                    f'{len(usuarios)} filas vacías debajo de los encabezados.')

    estilos = {col: copy(hoja.cell(row=inicio, column=col)._style) for col in todas}
    for i, usuario in enumerate(usuarios):
        for col in todas:
            celda = hoja.cell(row=inicio + i, column=col)
            if i and not celda.has_style:
                celda._style = copy(estilos[col])
            valor = i + 1 if col in columnas_numeracion else _valor(usuario, columnas[col])
            # openpyxl pone 'yyyy-mm-dd' al asignar una fecha: se respeta el formato
            # de la plantilla y, si no tenía, se usa el venezolano
            formato_plantilla = celda.number_format
            celda.value = valor
            if isinstance(valor, date):
                celda.number_format = formato_plantilla if formato_plantilla != 'General' else 'DD/MM/YYYY'


# ==========================================
# --- ORQUESTADOR ---
# ==========================================

def generar_reporte(archivo, instrucciones):
    """
    Devuelve (BytesIO con el Excel relleno, resumen para mostrar al usuario).
    Lanza ErrorGestor con un mensaje claro si algo impide generar el reporte.
    """
    try:
        libro = load_workbook(archivo)
    except Exception:
        raise ErrorGestor('No se pudo abrir el archivo. Verifica que sea un Excel .xlsx válido.')
    hoja = libro.worksheets[0]

    fila_encabezado, encabezados = detectar_encabezados(hoja)
    if not encabezados:
        raise ErrorGestor(f'No se encontraron encabezados en las primeras {FILAS_ENCABEZADO} filas '
                          f'de la hoja "{hoja.title}".')

    columnas_numeracion = [c for c, t in encabezados.items() if _normalizar(t) in ENCABEZADOS_NUMERACION]
    para_ia = list(dict.fromkeys(t for c, t in encabezados.items() if c not in columnas_numeracion))

    analisis = analizar_formato_ministerio(instrucciones, para_ia)
    if not analisis:
        raise ErrorGestor('El Cerebro (Gemini) no respondió. Puede ser la cuota de la API, la clave '
                          'o la conexión; intenta de nuevo en un minuto.')

    mapeo = analisis['mapeo_columnas']
    columnas = {c: mapeo[t] for c, t in encabezados.items() if t in mapeo}
    if not columnas:
        raise ErrorGestor('El Cerebro no pudo relacionar ningún encabezado de la plantilla con los '
                          'datos del personal. Revisa que la fila de encabezados sea la correcta.')

    filtros = analisis['filtros']
    usuarios = filtrar_trabajadores(filtros)
    if not usuarios:
        cargo = filtros['cargo_requerido']
        raise ErrorGestor('Ningún trabajador coincide con el filtro'
                          + (f' de cargo "{cargo}"' if cargo else '')
                          + (' (solo activos)' if filtros['solo_activos'] else '') + '.')

    inyectar(hoja, fila_encabezado, columnas, columnas_numeracion, usuarios)

    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)

    resumen = {
        'explicacion': filtros['explicacion'],
        'cargo_requerido': filtros['cargo_requerido'],
        'solo_activos': filtros['solo_activos'],
        'total': len(usuarios),
        'hoja': hoja.title,
        'fila_encabezado': fila_encabezado,
        'columnas': {encabezados[c]: atributo for c, atributo in columnas.items()},
        'sin_mapear': [t for c, t in encabezados.items() if c not in columnas and c not in columnas_numeracion],
        'numeracion': [encabezados[c] for c in columnas_numeracion],
    }
    return salida, resumen
